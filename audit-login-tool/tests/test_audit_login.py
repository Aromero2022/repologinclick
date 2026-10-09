"""Tests para audit_login.py — cubre parseo de JSON, plataforma, CSV truncado y vacío."""

import io
import json
import os
import tempfile

import pandas as pd
import pytest

sys_path = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
import sys

sys.path.insert(0, sys_path)

from audit_login import (
    device_type,
    extract_platform,
    parse_audit_context,
    read_csv_safe,
    process_files,
    write_excel,
    VALID_PLATFORMS,
)


# ---------------------------------------------------------------------------
# Helpers para generar CSV de prueba
# ---------------------------------------------------------------------------

HEADER = (
    "Operator ID,Operator Name,Proxy ID,Proxy Name,"
    "Secondary Login Operator ID,Secondary Login Operator Name,"
    "Audit Type,Timestamp,Operation Completed?,Correlation ID,Audit Context"
)


def _make_ctx(overrides: dict | None = None, *, replace: bool = False) -> str:
    base = {
        "Account ID": "TRUFUT",
        "Global User ID": "user001",
        "Login Name": "12345",
        "User Name": "67890",
        "Assignment ID": "99999",
        "IP Address": "10.0.0.1",
        "Login Channel": "WEB",
        "Login Method": "SSO",
        "Login Device": "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36",
        "Session ID": "sess-abc-123",
        "Action": "Login",
    }
    if overrides:
        if replace:
            base = overrides
        else:
            base.update(overrides)
    raw = json.dumps(base)
    return '"' + raw.replace('"', '""') + '"'


def _make_csv_line(ctx_overrides: dict | None = None, *, replace_ctx: bool = False, **kw) -> str:
    fields = {
        "op_id": "USR001",
        "op_name": "John Doe",
        "proxy_id": "",
        "proxy_name": "",
        "sec_id": "",
        "sec_name": "",
        "audit_type": "User Login",
        "timestamp": "2026-09-15T12:00:00.000Z",
        "completed": "TRUE",
        "corr_id": "corr-001",
    }
    fields.update(kw)
    ctx = _make_ctx(ctx_overrides, replace=replace_ctx)
    return (
        f"{fields['op_id']},{fields['op_name']},{fields['proxy_id']},"
        f"{fields['proxy_name']},{fields['sec_id']},{fields['sec_name']},"
        f"{fields['audit_type']},{fields['timestamp']},{fields['completed']},"
        f"{fields['corr_id']},{ctx}"
    )


def _write_csv(tmpdir, filename, lines):
    path = os.path.join(tmpdir, filename)
    content = HEADER + "\n" + "\n".join(lines) + "\n"
    with open(path, "w", encoding="utf-8-sig") as f:
        f.write(content)
    return path


# ---------------------------------------------------------------------------
# Tests de parse_audit_context
# ---------------------------------------------------------------------------


class TestParseAuditContext:
    def test_web_login(self):
        ctx = {
            "Account ID": "TRUFUT",
            "Login Channel": "WEB",
            "Login Method": "SSO",
            "Login Device": "Mozilla/5.0 (Windows NT 10.0)",
            "Action": "Login",
        }
        result = parse_audit_context(json.dumps(ctx))
        assert result["Login Channel"] == "WEB"
        assert result["Action"] == "Login"

    def test_mobile_login(self):
        device_json = json.dumps(
            {"appName": "SuccessFactors", "deviceName": "iPhone12,1", "osType": "iOS"}
        )
        ctx = {
            "Login Channel": "MOBILE",
            "Login Method": "TOKEN",
            "Login Device": device_json,
            "Action": "Login",
        }
        result = parse_audit_context(json.dumps(ctx))
        assert result["Login Channel"] == "MOBILE"
        assert result["Login Device"].startswith("{")

    def test_logout(self):
        ctx = {
            "Account ID": "TRUFUT",
            "Session ID": "sess-123",
            "Logout Method": "Time-Out",
            "Action": "Logout",
        }
        result = parse_audit_context(json.dumps(ctx))
        assert result["Action"] == "Logout"
        assert "Login Device" not in result

    def test_failed_login(self):
        ctx = {
            "IP Address": "192.168.1.1",
            "Login Channel": "WEB",
            "Failure Reason": "Invalid credentials",
        }
        result = parse_audit_context(json.dumps(ctx))
        assert "Action" not in result
        assert result["Failure Reason"] == "Invalid credentials"

    def test_empty_string(self):
        assert parse_audit_context("") == {}

    def test_invalid_json(self):
        assert parse_audit_context("not json") == {}


# ---------------------------------------------------------------------------
# Tests de extract_platform
# ---------------------------------------------------------------------------


class TestExtractPlatform:
    def test_windows(self):
        ua = "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36"
        assert extract_platform(ua) == "Windows"

    def test_android(self):
        ua = "Mozilla/5.0 (Android 16; Mobile; rv:128.0) Gecko/128.0 Firefox/128.0"
        assert extract_platform(ua) == "Android"

    def test_linux(self):
        ua = "Mozilla/5.0 (Linux; Android 10; K) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Mobile"
        assert extract_platform(ua) == "Linux"

    def test_linux_android_as_mobile(self):
        ua = "Mozilla/5.0 (Linux; Android 10; K) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0 Mobile"
        assert extract_platform(ua, android_as_mobile=True) == "Android"

    def test_iphone(self):
        ua = "Mozilla/5.0 (iPhone; CPU iPhone OS 17_0 like Mac OS X)"
        assert extract_platform(ua) == "iPhone"

    def test_ipad(self):
        ua = "Mozilla/5.0 (iPad; CPU OS 17_0 like Mac OS X)"
        assert extract_platform(ua) == "iPad"

    def test_macintosh(self):
        ua = "Mozilla/5.0 (Macintosh; Intel Mac OS X 14_0)"
        assert extract_platform(ua) == "Macintosh"

    def test_x11(self):
        ua = "Mozilla/5.0 (X11; Linux x86_64; rv:109.0) Gecko/20100101"
        assert extract_platform(ua) == "X11"

    def test_empty_string(self):
        assert extract_platform("") == "MOBILE"

    def test_json_device(self):
        assert extract_platform('{"appName":"SuccessFactors"}') == "MOBILE"

    def test_no_parens(self):
        assert extract_platform("SomeBot/1.0") == "MOBILE"


# ---------------------------------------------------------------------------
# Tests de device_type
# ---------------------------------------------------------------------------


class TestDeviceType:
    @pytest.mark.parametrize(
        "platform,expected",
        [
            ("Windows", "Laptop"),
            ("Linux", "Laptop"),
            ("Macintosh", "Laptop"),
            ("X11", "Laptop"),
            ("iPhone", "Mobile"),
            ("Android", "Mobile"),
            ("MOBILE", "Mobile"),
            ("iPad", "Tablet"),
        ],
    )
    def test_known_platforms(self, platform, expected):
        assert device_type(platform) == expected

    def test_unknown_platform(self):
        assert device_type("UnknownOS") == ""


# ---------------------------------------------------------------------------
# Tests de read_csv_safe
# ---------------------------------------------------------------------------


class TestReadCsvSafe:
    def test_normal_csv(self, tmp_path):
        line = _make_csv_line()
        path = _write_csv(str(tmp_path), "normal.csv", [line])
        df, warnings = read_csv_safe(path)
        assert len(df) == 1
        assert len(warnings) == 0

    def test_truncated_csv(self, tmp_path):
        line1 = _make_csv_line(timestamp="2026-09-01T06:00:00.000Z")
        line2 = _make_csv_line(timestamp="2026-09-02T06:00:00.000Z")
        truncated = line2[:50]  # cortar a media línea
        path = str(tmp_path / "truncated.csv")
        content = HEADER + "\n" + line1 + "\n" + truncated
        with open(path, "w", encoding="utf-8-sig") as f:
            f.write(content)

        df, warnings = read_csv_safe(path)
        assert len(df) == 1
        assert len(warnings) == 1
        assert "truncada" in warnings[0].lower()

    def test_empty_csv(self, tmp_path):
        path = str(tmp_path / "empty.csv")
        with open(path, "w", encoding="utf-8-sig") as f:
            f.write(HEADER + "\n")
        df, warnings = read_csv_safe(path)
        assert len(df) == 0
        assert len(warnings) == 0


# ---------------------------------------------------------------------------
# Tests de process_files (integración)
# ---------------------------------------------------------------------------


class TestProcessFiles:
    def test_web_login_end_to_end(self, tmp_path):
        line = _make_csv_line(
            {"Login Device": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"},
            timestamp="2026-09-15T12:00:00.000Z",
        )
        path = _write_csv(str(tmp_path), "test.csv", [line])
        df, infos, warnings = process_files([path])

        assert len(df) == 1
        assert df.iloc[0]["Plataforma"] == "Windows"
        assert df.iloc[0]["Tipo"] == "Laptop"
        assert df.iloc[0]["Login Device"] == "Mozilla/5.0 (Windows NT 10.0; Win64; x64)"

    def test_logout_becomes_mobile(self, tmp_path):
        logout_ctx = {
            "Account ID": "TRUFUT",
            "Global User ID": "user001",
            "Login Name": "12345",
            "User Name": "67890",
            "Assignment ID": "99999",
            "Session ID": "sess-logout-001",
            "Logout Method": "Time-Out",
            "Action": "Logout",
        }
        line = _make_csv_line(logout_ctx, replace_ctx=True, timestamp="2026-09-15T13:00:00.000Z")
        path = _write_csv(str(tmp_path), "logout.csv", [line])
        df, _, _ = process_files([path])

        assert len(df) == 1
        assert df.iloc[0]["Plataforma"] == "MOBILE"
        assert df.iloc[0]["Tipo"] == "Mobile"
        assert df.iloc[0]["Login Device"] == "MOBILE"
        assert df.iloc[0]["Session ID"] == "sess-logout-001"

    def test_mobile_app_json_device(self, tmp_path):
        device = json.dumps({"appName": "SuccessFactors", "deviceName": "iPhone12,1"})
        line = _make_csv_line(
            {"Login Channel": "MOBILE", "Login Method": "TOKEN", "Login Device": device},
            timestamp="2026-09-15T14:00:00.000Z",
        )
        path = _write_csv(str(tmp_path), "mobile.csv", [line])
        df, _, _ = process_files([path])

        assert df.iloc[0]["Plataforma"] == "MOBILE"
        assert df.iloc[0]["Login Device"] == "MOBILE"

    def test_failed_login(self, tmp_path):
        ctx = {
            "IP Address": "192.168.1.1",
            "Login Channel": "WEB",
            "Login Method": "SSO",
            "Login Device": "Mozilla/5.0 (Windows NT 10.0; Win64; x64)",
            "Failure Reason": "Invalid",
        }
        line = _make_csv_line(ctx, replace_ctx=True, completed="FALSE", timestamp="2026-09-15T15:00:00.000Z")
        path = _write_csv(str(tmp_path), "failed.csv", [line])
        df, _, _ = process_files([path])

        assert len(df) == 1
        assert df.iloc[0]["Operation Completed?"] == "FALSE"
        assert df.iloc[0]["Action"] == ""

    def test_dedup(self, tmp_path):
        line = _make_csv_line(timestamp="2026-09-15T12:00:00.000Z")
        path = _write_csv(str(tmp_path), "dup.csv", [line, line])
        df, _, warnings = process_files([path], dedup=True)
        assert len(df) == 1
        assert any("duplicados" in w.lower() for w in warnings)

    def test_solo_logins(self, tmp_path):
        login = _make_csv_line({"Action": "Login"}, timestamp="2026-09-15T12:00:00.000Z")
        logout_ctx = {
            "Account ID": "TRUFUT",
            "Session ID": "sess-001",
            "Logout Method": "Time-Out",
            "Action": "Logout",
        }
        logout = _make_csv_line(logout_ctx, replace_ctx=True, timestamp="2026-09-15T13:00:00.000Z")
        path = _write_csv(str(tmp_path), "mixed.csv", [login, logout])
        df, _, _ = process_files([path], solo_logins=True)
        assert len(df) == 1
        assert df.iloc[0]["Action"] == "Login"


# ---------------------------------------------------------------------------
# Tests de write_excel
# ---------------------------------------------------------------------------


class TestWriteExcel:
    def test_creates_valid_xlsx(self, tmp_path):
        line = _make_csv_line(timestamp="2026-09-15T12:00:00.000Z")
        csv_path = _write_csv(str(tmp_path), "test.csv", [line])
        df, _, _ = process_files([csv_path])
        xlsx_path = str(tmp_path / "output.xlsx")
        result_path = write_excel(df, xlsx_path)

        assert os.path.exists(result_path)
        assert os.path.getsize(result_path) > 0

    def test_formula_in_column_u(self, tmp_path):
        line = _make_csv_line(timestamp="2026-09-15T12:00:00.000Z")
        csv_path = _write_csv(str(tmp_path), "test.csv", [line])
        df, _, _ = process_files([csv_path])
        xlsx_path = str(tmp_path / "formula.xlsx")
        write_excel(df, xlsx_path)

        import openpyxl

        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.active
        assert ws.cell(1, 21).value == "Tipo Dispositivo"
        cell_u2 = ws.cell(2, 21).value
        assert cell_u2 is not None
        assert "_xlfn.IFS" in str(cell_u2)
        wb.close()

    def test_21_columns(self, tmp_path):
        line = _make_csv_line(timestamp="2026-09-15T12:00:00.000Z")
        csv_path = _write_csv(str(tmp_path), "test.csv", [line])
        df, _, _ = process_files([csv_path])
        xlsx_path = str(tmp_path / "cols.xlsx")
        write_excel(df, xlsx_path)

        import openpyxl

        wb = openpyxl.load_workbook(xlsx_path)
        ws = wb.active
        headers = [ws.cell(1, j).value for j in range(1, 22)]
        assert len(headers) == 21
        assert headers[0] == "Operator ID"
        assert headers[19] == "Plataforma"
        assert headers[20] == "Tipo Dispositivo"
        wb.close()


# ---------------------------------------------------------------------------
# Tests de plataformas válidas
# ---------------------------------------------------------------------------


class TestPlatformValidation:
    def test_all_known_platforms_have_device_type(self):
        for p in VALID_PLATFORMS:
            assert device_type(p) != "", f"Plataforma {p} no tiene tipo de dispositivo"
